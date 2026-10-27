"""El çizimi takip dosyası: docs/el_cizimi/el_cizimi_takip.xlsx.

Envanterden üretilir. Yeniden üretilince çizerin doldurduğu sütunlar (Durum,
tarihler, harcanan saat, notlar) kimliğe göre eski dosyadan taşınır; envanterden
düşen bir satırın girdisi kaybolmaz, "Eski satırlar" sayfasına yazılır.

    python3 -I tools/el_cizimi/takip_excel.py [--oyunda ID ...]
`--oyunda`: alım aracı bir çizimi oyuna işlediğinde durumu "Oyunda" yapar.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import envanter  # noqa: E402
from sablon_uret import ZIP_GROUPS  # noqa: E402

OUT = envanter.ROOT / "docs" / "el_cizimi" / "el_cizimi_takip.xlsx"
STATUSES = ["Başlanmadı", "Çizgi", "Renk", "Bitti", "Yüklendi", "Oyunda", "Gerek yok"]
DONE = ["Bitti", "Yüklendi", "Oyunda"]
STATUS_FILL = {"Çizgi": "FFF2CC", "Renk": "FCE4D6", "Bitti": "DDEBF7", "Yüklendi": "E2EFDA",
               "Oyunda": "A9D08E", "Gerek yok": "D9D9D9"}

COLUMNS = [
    ("Kimlik", 24, "id"), ("Kategori", 11, "kategori"), ("Alt grup", 20, "alt_grup"), ("Ad", 30, "ad"),
    ("Ne çizilecek", 50, "ne_cizilecek"), ("Renklendirme", 22, "renk"), ("Tuval (px)", 14, "tuval"),
    ("Katmanlar (alttan üste)", 34, "katmanlar"), ("Teslim", 18, "teslim"), ("Şablon paketi", 14, "zip"),
    ("Bugünkü kaynak", 30, "kaynak"), ("Öncelik", 8, "oncelik"), ("Tahmini saat", 10, "saat"),
    ("Önce bitmesi gereken", 20, "bagimlilik"), ("Koda bağlantı gerekir", 10, "kod"),
    ("Oyundaki dosya", 44, "hedef"), ("Not", 36, "not"),
]
USER_COLUMNS = [("Durum", 13), ("Başlama", 11), ("Bitiş", 11), ("Harcanan saat", 10), ("Notlarım", 36)]

ARIAL = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="3A3226")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")


def _previous() -> tuple[dict, list]:
    if not OUT.exists():
        return {}, []
    wb = load_workbook(OUT)
    ws = wb["Liste"]
    head = [c.value for c in ws[1]]
    idx = {name: head.index(name) for name, _ in USER_COLUMNS if name in head}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        out[row[0]] = {name: row[i] for name, i in idx.items()}
    old = []
    if "Eski satırlar" in wb.sheetnames:
        old = [list(r) for r in wb["Eski satırlar"].iter_rows(min_row=2, values_only=True)]
    return out, old


def build(in_game: list[str]) -> None:
    rows = envanter.all_rows()
    prev, old_rows = _previous()
    ids = {r["id"] for r in rows}
    for k, v in prev.items():
        if k not in ids and any(x not in (None, "", "Başlanmadı") for x in v.values()):
            old_rows.append([k] + [v.get(n) for n, _ in USER_COLUMNS])
    order = {c: i for i, c in enumerate(envanter.ORDER)}
    rows.sort(key=lambda r: (order.get(r["kategori"], 99), r["oncelik"], r["id"]))

    wb = Workbook()
    _how(wb.active, len(rows))
    ws = wb.create_sheet("Liste")
    _list(ws, rows, prev, set(in_game))
    _summary(wb.create_sheet("Özet", 0), len(rows))
    _layers(wb.create_sheet("Katman adları"))
    if old_rows:
        o = wb.create_sheet("Eski satırlar")
        o.append(["Kimlik"] + [n for n, _ in USER_COLUMNS])
        for r in old_rows:
            o.append(r)
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                if c.font is None or c.font.name != ARIAL:
                    c.font = Font(name=ARIAL, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                                  color=c.font.color if c.font else None)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("yazıldı:", OUT.relative_to(envanter.ROOT), len(rows), "satır")


def _head(ws, names_widths) -> None:
    for i, (name, width) in enumerate(names_widths, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def _list(ws, rows, prev, in_game) -> None:
    cols = [(n, w) for n, w, _ in COLUMNS] + USER_COLUMNS
    _head(ws, cols)
    first_user = len(COLUMNS) + 1
    for ri, r in enumerate(rows, start=2):
        values = {
            "zip": ZIP_GROUPS[r["kategori"]] + ".zip", "katmanlar": ", ".join(r["katmanlar"]),
            "kod": "Evet" if r["kod_gerekir"] else "",
        }
        for ci, (_, _, key) in enumerate(COLUMNS, start=1):
            v = values.get(key, r.get(key))
            c = ws.cell(row=ri, column=ci, value=v)
            c.alignment = Alignment(wrap_text=key in ("ne_cizilecek", "not", "katmanlar", "ad"), vertical="top")
            c.font = Font(name=ARIAL, size=9)
        p = prev.get(r["id"], {})
        status = "Oyunda" if r["id"] in in_game else (p.get("Durum") or "Başlanmadı")
        user_vals = [status, p.get("Başlama"), p.get("Bitiş"), p.get("Harcanan saat"), p.get("Notlarım")]
        for k, v in enumerate(user_vals):
            c = ws.cell(row=ri, column=first_user + k, value=v)
            c.font = Font(name=ARIAL, size=9)
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if k in (1, 2):
                c.number_format = "d.m.yyyy"
    last = len(rows) + 1
    status_col = get_column_letter(first_user)
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(STATUSES), allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("%s2:%s%d" % (status_col, status_col, last))
    for k in (1, 2):
        col = get_column_letter(first_user + k)
        d = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True)
        ws.add_data_validation(d)
        d.add("%s2:%s%d" % (col, col, last))
    full = "A2:%s%d" % (get_column_letter(len(cols)), last)
    for status, color in STATUS_FILL.items():
        ws.conditional_formatting.add(full, FormulaRule(formula=['$%s2="%s"' % (status_col, status)],
                                                        fill=PatternFill("solid", fgColor=color)))
    for k in range(len(USER_COLUMNS)):
        ws.cell(row=1, column=first_user + k).fill = PatternFill("solid", fgColor="8A6D1E")
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), last)


def _summary(ws, n: int) -> None:
    last = n + 1
    s_col = get_column_letter(len(COLUMNS) + 1)
    h_col = get_column_letter(len(COLUMNS) + 4)
    L = "Liste!"
    rng = lambda col: "%s$%s$2:$%s$%d" % (L, col, col, last)  # noqa: E731
    ws["A1"] = "Wayborne el çizimi - ilerleme"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
    ws["A2"] = "Bu sayfa formül: Liste sayfasındaki Durum ve Harcanan saat sütunlarından hesaplanır. Elle doldurma."
    ws["A2"].font = Font(name=ARIAL, italic=True, size=9, color="666666")
    heads = ["Kategori", "Toplam", "Bitti (Bitti+Yüklendi+Oyunda)", "Oyunda", "Gerek yok", "İlerleme",
             "Tahmini saat", "Kalan tahmini saat", "Harcanan saat", "AI kaynaklı kalan"]
    widths = [18, 9, 16, 9, 9, 10, 12, 14, 12, 14]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 30
    cats = envanter.ORDER
    for i, cat in enumerate(cats, start=5):
        a = "$A%d" % i
        ws.cell(row=i, column=1, value=cat)
        ws.cell(row=i, column=2, value="=COUNTIF(%s,%s)" % (rng("B"), a))
        ws.cell(row=i, column=3, value="=" + "+".join(
            'COUNTIFS(%s,%s,%s,"%s")' % (rng("B"), a, rng(s_col), s) for s in DONE))
        ws.cell(row=i, column=4, value='=COUNTIFS(%s,%s,%s,"Oyunda")' % (rng("B"), a, rng(s_col)))
        ws.cell(row=i, column=5, value='=COUNTIFS(%s,%s,%s,"Gerek yok")' % (rng("B"), a, rng(s_col)))
        ws.cell(row=i, column=6, value="=IF(B%d-E%d=0,0,C%d/(B%d-E%d))" % (i, i, i, i, i))
        ws.cell(row=i, column=7, value="=SUMIF(%s,%s,%s)" % (rng("B"), a, rng("M")))
        ws.cell(row=i, column=8, value="=G%d-" % i + "-".join(
            'SUMIFS(%s,%s,%s,%s,"%s")' % (rng("M"), rng("B"), a, rng(s_col), s) for s in DONE + ["Gerek yok"]))
        ws.cell(row=i, column=9, value="=SUMIF(%s,%s,%s)" % (rng("B"), a, rng(h_col)))
        ws.cell(row=i, column=10, value='=COUNTIFS(%s,%s,%s,"AI*")' % (rng("B"), a, rng("K")) + "".join(
            '-COUNTIFS(%s,%s,%s,"AI*",%s,"%s")' % (rng("B"), a, rng("K"), rng(s_col), s) for s in DONE + ["Gerek yok"]))
    t = len(cats) + 5
    ws.cell(row=t, column=1, value="TOPLAM").font = Font(name=ARIAL, bold=True)
    for col in range(2, 11):
        letter = get_column_letter(col)
        if col == 6:
            ws.cell(row=t, column=col, value="=IF(B%d-E%d=0,0,C%d/(B%d-E%d))" % (t, t, t, t, t))
        else:
            ws.cell(row=t, column=col, value="=SUM(%s5:%s%d)" % (letter, letter, t - 1))
        ws.cell(row=t, column=col).font = Font(name=ARIAL, bold=True)
    for r in range(5, t + 1):
        ws.cell(row=r, column=6).number_format = "0.0%"
        for col in (7, 8, 9):
            ws.cell(row=r, column=col).number_format = "#,##0"
    p = t + 2
    ws.cell(row=p, column=1, value="Önceliğe göre").font = Font(name=ARIAL, bold=True)
    labels = {1: "1 - AI görseller (beyan kalkar)", 2: "2 - En çok görünen render'lar", 3: "3 - Geri kalan ve yeni"}
    for k, pr in enumerate((1, 2, 3)):
        r = p + 1 + k
        ws.cell(row=r, column=1, value=labels[pr])
        ws.cell(row=r, column=2, value="=COUNTIF(%s,%d)" % (rng("L"), pr))
        ws.cell(row=r, column=3, value="=" + "+".join(
            'COUNTIFS(%s,%d,%s,"%s")' % (rng("L"), pr, rng(s_col), s) for s in DONE))
        ws.cell(row=r, column=6, value="=IF(B%d=0,0,C%d/B%d)" % (r, r, r)).number_format = "0.0%"
        ws.cell(row=r, column=7, value="=SUMIF(%s,%d,%s)" % (rng("L"), pr, rng("M"))).number_format = "#,##0"
    ws.column_dimensions["A"].width = 34
    ws.cell(row=p + 5, column=1, value="Tahmini saatler kaba bir plan içindir (bkz. Nasıl sayfası); "
                                       "gerçek süreyi Harcanan saat sütununa yazdıkça karşılaştırılır.").font = \
        Font(name=ARIAL, italic=True, size=9, color="666666")


def _how(ws, n: int) -> None:
    ws.title = "Nasıl"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    lines = [
        ("Bu dosya", "Oyundaki ve oyunda olması gereken %d görselin listesi. Süreç belgesi: docs/el_cizimi/SUREC.md." % n),
        ("Senin dolduracağın", "Liste sayfasında sarı başlıklı beş sütun: Durum, Başlama, Bitiş, Harcanan saat, Notlarım. "
                               "Gerisini araç yazar; elle değiştirme, yeniden üretilince geri gelir."),
        ("Durum değerleri", "Başlanmadı → Çizgi (lineart bitti) → Renk (düz renk sürüyor) → Bitti (tablette tamam) → "
                            "Yüklendi (GitHub'a yüklendi) → Oyunda (Claude işledi, oyunda görünüyor). "
                            "Gerek yok: bu görseli çizmemeye karar verdin."),
        ("Oyunda", "Bu durumu Claude yazar: yüklediğin dosyayı doğrulayıp oyuna işledikten ve ekran görüntüsüyle "
                   "kontrol ettikten sonra."),
        ("Öncelik", "1: bugün yapay zekâ üretimi olan görseller (bitince AI beyanı kalkar). 2: oyunda en çok görünen "
                    "render'lar (beden, yaygın giysiler, hayvanlar, vagon, kasabaların gündüz resmi, temel yol "
                    "katmanları). 3: geri kalan ve henüz olmayan parçalar."),
        ("Tahmini saat", "Kaba bir plan: bir ikon 0,75 saat, bir giysi 2,5, bir beden 8, bir kasaba gündüz resmi 10. "
                         "Kendi hızına göre değişir; Harcanan saat sütunu gerçek ortalamanı gösterir."),
        ("Önce bitmesi gereken", "Bu satır o satıra dayanıyor (giysi, bedenin çizgilerine oturur; bir yol "
                                 "varyantı ana biyomun çizimini yeniden renklendirir)."),
        ("Teslim", "PSD (adlı katmanlar): ibis Paint'ten PSD olarak dışa aktar, katman adlarını şablondaki gibi bırak. "
                   "PNG (saydam): tek katman, saydam arka plan. Dosya adı Kimlik sütunuyla birebir aynı."),
        ("Renklendirme", "Gri tonlarla: oyun rengi kendisi çarpar (ten, saç rengi, temel takım). Tam renk: "
                         "çizdiğin renk oyunda aynen görünür."),
        ("Koda bağlantı gerekir", "Evet: oyunda henüz böyle bir resim yok; teslim edilince Claude oyuna bağlar."),
    ]
    ws["A1"] = "Nasıl kullanılır"
    ws["A1"].font = Font(name=ARIAL, bold=True, size=14)
    for i, (k, v) in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name=ARIAL, bold=True, size=10)
        c = ws.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 42
    r = len(lines) + 5
    ws.cell(row=r, column=1, value="Örnek satır").font = Font(name=ARIAL, bold=True, size=10)
    example = [("Kimlik", "KYF-peasant_shirt-E"), ("Durum", "Renk"), ("Başlama", "03.11.2026"),
               ("Bitiş", ""), ("Harcanan saat", "1,5"), ("Notlarım", "Yaka çizgisi tamam, gölge kaldı")]
    for k, (a, b) in enumerate(example):
        ws.cell(row=r + 1 + k, column=1, value=a).font = Font(name=ARIAL, size=10)
        c = ws.cell(row=r + 1 + k, column=2, value=b)
        c.font = Font(name=ARIAL, size=10, color="0000FF")
        if a != "Kimlik":
            c.fill = INPUT_FILL


def _layers(ws) -> None:
    ws.append(["Tür", "Katman adları (PSD'de alttan üste)", "Açıklama"])
    for c in ws[1]:
        c.font = Font(name=ARIAL, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    ws.append(["Figür (beden, giysi, saç)", ", ".join(envanter.FIG_LAYERS),
               "arka_kol: uzak kol · arka_bacak: uzak bacak · govde: gövde ve baş · on_bacak: yakın bacak · "
               "on_kol: yakın kol. Bir parça birden çok katmana düşebilir (gömleğin kolları kol katmanına)."])
    ws.append(["Hayvan", ", ".join(envanter.BEAST_LAYERS), "Uzak bacaklar gövdenin arkasında; kuyruk ayrı sallanır."])
    ws.append(["Brandalı vagon", "govde, yuk, branda, teker_arka, teker_on",
               "Teker kendi katmanında döner; branda tabanından salınır."])
    ws.append(["Kapı", "kasaba, govde", "kasaba: surun arkasındaki şehir; govde: sur ve kapı çerçevesi."])
    ws.append(["Değirmen", "govde, kanatlar", "Kanatlar kendi katmanında döner."])
    ws.append(["Birden çok çeşit (ağaç, çalı...)", "agac_1, agac_2... / cali_1...",
               "Her çeşit kendi katmanında, şablondaki kendi kutusunda."])
    ws.append(["Tek parça", "cizim", "PNG olarak teslim edilebilir; katman adı gerekmez."])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 90
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")


if __name__ == "__main__":
    args = sys.argv[1:]
    marks = args[args.index("--oyunda") + 1:] if "--oyunda" in args else []
    build(marks)
